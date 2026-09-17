"""Base de donnees complete des bagages - filtres avances + export Excel."""
from __future__ import annotations

import os
from datetime import date, datetime, timedelta
from pathlib import Path

from PySide6.QtCore import Qt, QDate
from PySide6.QtGui import QColor, QFont
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit,
    QComboBox, QPushButton, QTableWidget, QTableWidgetItem,
    QHeaderView, QDateEdit, QMessageBox, QFileDialog, QApplication,
)
from sqlalchemy.orm import joinedload

from database.session import get_session
from models.luggage import Luggage
from resources import theme as T
from services.agency_context import current_agency_id
from utils.formatters import format_fc
from views.admin.widgets import style_table, secondary_btn
from views.widgets.loading import LoadingOverlay, TableSkeleton


def _q_to_date(qd):
    return date(qd.year(), qd.month(), qd.day())


def _date_to_q(d):
    return QDate(d.year, d.month, d.day)


STATUTS = {
    "Tous": None,
    "Enregistre": "enregistre",
    "Charge": "charge",
    "Livre": "livre",
    "Annule": "annule",
}

STATUTS_FR = {
    "enregistre": "Enregistre",
    "charge": "Charge",
    "livre": "Livre",
    "annule": "Annule",
}

STATUT_COLORS = {
    "enregistre": ("#1a73e8", "#e8f0fe"),
    "charge":     ("#f09300", "#fff3e0"),
    "livre":      ("#1e8c45", "#e6f4ea"),
    "annule":     ("#c62828", "#fce8e6"),
}

COLUMNS = [
    "Numero", "Expediteur", "Destinataire", "Tel. Dest.",
    "Poids (kg)", "Total (FC)", "Trajet", "Bus", "Statut", "Date",
]


class BagagesDBView(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self._data = []
        self._build_ui()
        self._loading_overlay = LoadingOverlay(self, "Chargement de la base bagages…")
        self._table_skeleton = TableSkeleton(self.table, rows=8)

    def _field_style(self):
        return (
            "QLineEdit, QDateEdit, QComboBox {"
            "  background: #1e1e2e; color: #e0e0e0;"
            "  border: 1px solid #333; border-radius: 6px;"
            "  padding: 4px 8px; font-size:13px; }"
            "QLineEdit:focus, QDateEdit:focus { border: 1px solid #1a73e8; }"
        )

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(10)

        title_row = QHBoxLayout()
        title = QLabel("Base de donnees Bagages")
        title.setStyleSheet(f"font-size:18px; font-weight:700; color:{T.TEXT_PRIMARY};")
        title_row.addWidget(title)
        title_row.addStretch()
        self.lbl_count = QLabel("0 bagages")
        self.lbl_count.setStyleSheet(f"color:{T.TEXT_SECONDARY}; font-size:13px;")
        title_row.addWidget(self.lbl_count)
        root.addLayout(title_row)

        row1 = QHBoxLayout()
        row1.setSpacing(8)
        self.search_code = QLineEdit()
        self.search_code.setPlaceholderText("Code (NG-260914-...)")
        self.search_code.setFixedHeight(34)
        self.search_code.textChanged.connect(self._apply_filters)
        self.search_code.setStyleSheet(self._field_style())
        row1.addWidget(self.search_code, 2)
        self.search_name = QLineEdit()
        self.search_name.setPlaceholderText("Expediteur / Destinataire")
        self.search_name.setFixedHeight(34)
        self.search_name.textChanged.connect(self._apply_filters)
        self.search_name.setStyleSheet(self._field_style())
        row1.addWidget(self.search_name, 2)
        self.search_phone = QLineEdit()
        self.search_phone.setPlaceholderText("Telephone")
        self.search_phone.setFixedHeight(34)
        self.search_phone.textChanged.connect(self._apply_filters)
        self.search_phone.setStyleSheet(self._field_style())
        row1.addWidget(self.search_phone, 1)
        root.addLayout(row1)

        row2 = QHBoxLayout()
        row2.setSpacing(8)
        lbl_du = QLabel("Du")
        lbl_du.setStyleSheet(f"color:{T.TEXT_SECONDARY};")
        row2.addWidget(lbl_du)
        self.date_from = QDateEdit()
        self.date_from.setCalendarPopup(True)
        self.date_from.setDate(_date_to_q(date.today() - timedelta(days=30)))
        self.date_from.setFixedHeight(34)
        self.date_from.setStyleSheet(self._field_style())
        self.date_from.dateChanged.connect(self._apply_filters)
        row2.addWidget(self.date_from)
        lbl_au = QLabel("au")
        lbl_au.setStyleSheet(f"color:{T.TEXT_SECONDARY};")
        row2.addWidget(lbl_au)
        self.date_to = QDateEdit()
        self.date_to.setCalendarPopup(True)
        self.date_to.setDate(_date_to_q(date.today()))
        self.date_to.setFixedHeight(34)
        self.date_to.setStyleSheet(self._field_style())
        self.date_to.dateChanged.connect(self._apply_filters)
        row2.addWidget(self.date_to)
        row2.addSpacing(10)
        lbl_s = QLabel("Statut")
        lbl_s.setStyleSheet(f"color:{T.TEXT_SECONDARY};")
        row2.addWidget(lbl_s)
        self.cmb_statut = QComboBox()
        for label, val in STATUTS.items():
            self.cmb_statut.addItem(label, val)
        self.cmb_statut.setFixedHeight(34)
        self.cmb_statut.setStyleSheet(self._field_style())
        self.cmb_statut.currentIndexChanged.connect(self._apply_filters)
        row2.addWidget(self.cmb_statut)
        row2.addStretch()
        btn_refresh = secondary_btn("Actualiser")
        btn_refresh.setFixedHeight(34)
        btn_refresh.clicked.connect(self.refresh)
        row2.addWidget(btn_refresh)
        self.btn_excel = QPushButton("Exporter Excel")
        self.btn_excel.setObjectName("primaryBtn")
        self.btn_excel.setFixedHeight(34)
        self.btn_excel.setCursor(Qt.CursorShape.PointingHandCursor)
        self.btn_excel.clicked.connect(self._export_excel)
        row2.addWidget(self.btn_excel)
        root.addLayout(row2)

        self.table = QTableWidget(0, len(COLUMNS))
        self.table.setHorizontalHeaderLabels(COLUMNS)
        style_table(self.table)
        self.table.setSortingEnabled(True)
        hh = self.table.horizontalHeader()
        hh.setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        for i, w in enumerate([155, 150, 150, 110, 80, 95, 150, 70, 90, 140]):
            self.table.setColumnWidth(i, w)
        root.addWidget(self.table, 1)

    def refresh(self):
        self._loading_overlay.show_loading("Chargement de la base bagages…")
        self._table_skeleton.show()
        QApplication.processEvents()
        try:
            session = get_session()
            try:
                aid = current_agency_id()
                q = (
                    session.query(Luggage)
                    .options(joinedload(Luggage.route), joinedload(Luggage.bus))
                    .order_by(Luggage.created_at.desc())
                )
                if aid is not None:
                    q = q.filter(Luggage.agency_id == aid)
                items = q.limit(5000).all()
                self._data = []
                for item in items:
                    self._data.append({
                        "id": item.id,
                        "numero": item.numero,
                        "sender_name": item.sender_name,
                        "sender_phone": item.sender_phone or "",
                        "recipient_name": item.recipient_name,
                        "recipient_phone": item.recipient_phone or "",
                        "poids": float(item.poids),
                        "total": item.total,
                        "statut": item.statut,
                        "route_label": item.route.short_label if item.route else "",
                        "bus_code": item.bus.code if item.bus else "-",
                        "created_at": item.created_at,
                    })
            finally:
                session.close()
            self._apply_filters()
        finally:
            self._table_skeleton.hide()
            self._loading_overlay.hide_loading()

    def _apply_filters(self):
        code = self.search_code.text().strip().lower()
        name = self.search_name.text().strip().lower()
        phone = self.search_phone.text().strip()
        statut = self.cmb_statut.currentData()
        d_from = _q_to_date(self.date_from.date())
        d_to = _q_to_date(self.date_to.date())

        filtered = []
        for item in self._data:
            item_date = item["created_at"].date() if item["created_at"] else date.today()
            if item_date < d_from or item_date > d_to:
                continue
            if statut and item["statut"] != statut:
                continue
            if code and code not in item["numero"].lower():
                continue
            if name and name not in item["sender_name"].lower() and name not in item["recipient_name"].lower():
                continue
            if phone and phone not in item["sender_phone"] and phone not in item["recipient_phone"]:
                continue
            filtered.append(item)
        self._populate_table(filtered)

    def _populate_table(self, items):
        self.table.setSortingEnabled(False)
        self.table.setRowCount(0)
        for item in items:
            row = self.table.rowCount()
            self.table.insertRow(row)
            statut_fr = STATUTS_FR.get(item["statut"], item["statut"])
            date_str = item["created_at"].strftime("%d/%m/%Y %H:%M") if item["created_at"] else ""
            vals = [
                item["numero"],
                item["sender_name"],
                item["recipient_name"],
                item["recipient_phone"] or "-",
                f"{item['poids']:.1f}",
                format_fc(item["total"]),
                item["route_label"],
                item["bus_code"],
                statut_fr,
                date_str,
            ]
            for col, v in enumerate(vals):
                cell = QTableWidgetItem(str(v))
                cell.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                if col == 0:
                    f = QFont()
                    f.setBold(True)
                    f.setFamily("Courier New")
                    cell.setFont(f)
                if col == 8:
                    fg, bg = STATUT_COLORS.get(item["statut"], ("#333", "#fff"))
                    cell.setForeground(QColor(fg))
                    cell.setBackground(QColor(bg))
                self.table.setItem(row, col, cell)
        self.table.setSortingEnabled(True)
        total = len(items)
        self.lbl_count.setText(f"<b>{total}</b> bagages trouves")

    def _export_excel(self):
        default_name = f"bagages_{date.today().strftime('%Y%m%d')}.xlsx"
        path, _ = QFileDialog.getSaveFileName(
            self, "Enregistrer Excel",
            str(Path.home() / "Downloads" / default_name),
            "Fichiers Excel (*.xlsx)",
        )
        if not path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = "Bagages"
            header_fill = PatternFill("solid", fgColor="1A73E8")
            header_font = Font(bold=True, color="FFFFFF", size=11)
            thin = Side(style="thin", color="CCCCCC")
            border = Border(left=thin, right=thin, bottom=thin)
            ws.append(COLUMNS)
            for col_idx in range(1, len(COLUMNS) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            statut_excel_fills = {
                "Enregistre": PatternFill("solid", fgColor="E8F0FE"),
                "Charge":     PatternFill("solid", fgColor="FFF3E0"),
                "Livre":      PatternFill("solid", fgColor="E6F4EA"),
                "Annule":     PatternFill("solid", fgColor="FCE8E6"),
            }
            for r in range(self.table.rowCount()):
                row_data = [
                    self.table.item(r, c).text() if self.table.item(r, c) else ""
                    for c in range(self.table.columnCount())
                ]
                ws.append(row_data)
                statut_val = row_data[8] if len(row_data) > 8 else ""
                fill = statut_excel_fills.get(statut_val)
                for col_idx in range(1, len(COLUMNS) + 1):
                    cell = ws.cell(row=r + 2, column=col_idx)
                    cell.border = border
                    cell.alignment = Alignment(vertical="center")
                    if fill and col_idx == 9:
                        cell.fill = fill
            for i, w in enumerate([22, 22, 22, 16, 11, 14, 22, 10, 13, 20], start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = "A2"
            ws.append([])
            ws.append(["Total bagages:", self.table.rowCount()])
            ws.append(["Exporte le:", datetime.now().strftime("%d/%m/%Y %H:%M")])
            wb.save(path)

            reply = QMessageBox.question(
                self, "Export reussi",
                f"Fichier enregistre :\n{path}\n\nOuvrir maintenant ?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            )
            if reply == QMessageBox.StandardButton.Yes:
                os.startfile(path)
        except ImportError:
            QMessageBox.critical(self, "Module manquant", "Installez openpyxl: pip install openpyxl")
        except Exception as e:
            QMessageBox.critical(self, "Erreur export", str(e))
